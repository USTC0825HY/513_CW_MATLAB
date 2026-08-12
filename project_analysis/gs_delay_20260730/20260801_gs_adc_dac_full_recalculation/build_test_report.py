from pathlib import Path
import json
import csv
import math

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BUNDLE = Path(__file__).resolve().parent
PROJECT = BUNDLE.parent.parent
FIG = BUNDLE / "Report_Figures"
TEMPLATE = (
    PROJECT
    / "01_Documents_文档"
    / "02_Test_Reports_测试报告"
    / "20260730_GS_延迟驱动.docx"
)
OUTPUT_DIR = PROJECT / "01_Documents_文档" / "02_Test_Reports_测试报告"
OUTPUT_DOCX = OUTPUT_DIR / "GS延迟驱动板_ADC_DAC测试数据分析报告_V1.0.docx"


def read_csv(path):
    with Path(path).open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def number(row, key):
    return float(row[key])


def truth(value):
    return str(value).strip().lower() in {"1", "true", "yes"}


def clear_document_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size=10.5, bold=False, color=None, name="宋体"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(document, text, bold_prefix=None):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(3)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_heading(document, text, level):
    style = {1: "一级标题", 2: "二级标题", 3: "三级标题"}[level]
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 15, 2: 13, 3: 11}[level], bold=True, name="黑体")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_properties.append(header)


def format_cell(cell, size=8.5, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_table(document, headers, rows, widths=None, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "报告表格样式" if "报告表格样式" in [s.name for s in document.styles] else "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        set_cell_shading(cell, "4472C4")
        format_cell(cell, size=font_size, bold=True, color=(255, 255, 255))
        if widths:
            cell.width = Cm(widths[index])
    set_repeat_table_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = "" if value is None else str(value)
            if widths:
                cells[index].width = Cm(widths[index])
            format_cell(
                cells[index],
                size=font_size,
                align=WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER,
            )
    document.add_paragraph()
    return table


def add_caption(document, text):
    paragraph = document.add_paragraph(style="图表题注")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=False)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_picture(document, path, caption, width_cm=16.2):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    with Path(path).open("rb") as image_stream:
        run = paragraph.add_run()
        run.add_picture(image_stream, width=Cm(width_cm))
    add_caption(document, caption)


def add_toc(document):
    paragraph = document.add_paragraph(style="目录标题")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("目录")
    set_run_font(run, size=16, bold=True, name="黑体")
    toc_paragraph = document.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run = OxmlElement("w:r")
    begin_run.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    instruction_run = OxmlElement("w:r")
    instruction_run.append(instruction)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run = OxmlElement("w:r")
    separate_run.append(separate)
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "打开文档后更新目录"
    text_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run = OxmlElement("w:r")
    end_run.append(end)
    toc_paragraph._p.extend(
        [begin_run, instruction_run, separate_run, text_run, end_run]
    )


def add_page_break(document):
    document.add_page_break()


def status_text(value):
    return "满足" if bool(value) else "不满足"


def build_report():
    sfdr = read_csv(BUNDLE / "SFDR" / "SFDR_all_channels_summary.csv")
    bandwidth = read_csv(BUNDLE / "Bandwidth" / "ADC_bandwidth_summary.csv")
    isolation = read_csv(BUNDLE / "Isolation" / "ADC_isolation_summary.csv")
    power = read_csv(BUNDLE / "PowerScale" / "ADC_power_scale_summary.csv")
    linearity = read_csv(BUNDLE / "AD766_Linearity" / "AD766_linearity_summary.csv")
    noise = read_csv(BUNDLE / "AD766_Noise" / "AD766_noise_summary.csv")
    anomalies = read_csv(BUNDLE / "anomaly_summary.csv")
    manifest = read_csv(BUNDLE / "source_manifest.csv")
    with (BUNDLE / "analysis_summary.json").open("r", encoding="utf-8") as stream:
        summary = json.load(stream)

    document = Document(str(TEMPLATE))
    clear_document_body(document)
    document.core_properties.title = "GS延迟驱动板 AD9245/AD766测试数据分析报告"
    document.core_properties.subject = "AD9245与AD766测试数据全量复算"
    document.core_properties.author = "测试数据复算与报告整理"
    document.core_properties.comments = "数据来源、算法参数及文件哈希见报告附录和结果包。"

    # Request automatic field refresh when Word opens the document.
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    title = document.add_paragraph(style="报告题名")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("GS延迟驱动板\nAD9245/AD766测试数据分析报告")
    set_run_font(title_run, size=22, bold=True, name="黑体")
    short_title = document.add_paragraph(style="报告短题名")
    short_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    short_title_run = short_title.add_run("GS延迟驱动板 ADC/DAC测试")
    set_run_font(short_title_run, size=10.5, name="宋体")
    document.add_paragraph()

    metadata_rows = [
        ["测试批次", "2026-07-30 GS延迟驱动板测试"],
        ["数据复算日期", "2026-08-01"],
        ["测试地点", "513（沿用原测试记录）"],
        ["测试人员", "赵庆、胡钰（沿用原测试记录）"],
        ["分析软件", "MATLAB R2025b"],
        ["数据范围", "AD9245四项测试及AD766线性/噪声，共49个原始CSV"],
        ["报告版本", "V1.0"],
        ["编写/审核/批准", "待填写 / 未记录 / 未记录"],
    ]
    add_table(document, ["项目", "内容"], metadata_rows, widths=[4.0, 12.0], font_size=10)

    add_caption(document, "表 1 修订记录")
    add_table(
        document,
        ["版本", "日期", "修订内容"],
        [["V1.0", "2026-08-01", "首次完成AD9245/AD766原始数据全量复算和报告整理"]],
        widths=[2.5, 3.5, 10.0],
        font_size=9,
    )

    add_page_break(document)
    add_toc(document)
    add_page_break(document)

    add_heading(document, "测试目的与结论摘要", 1)
    add_heading(document, "测试目的与范围", 2)
    add_text(
        document,
        "本报告依据《光梳频率梳延迟驱动板 测试细则（公开）》整理2026-07-30测试数据，"
        "对AD9245四路中速AD及AD766相关记录进行复算。原DATA_GS目录已在工程整理时迁移，"
        "本次实际数据源为02_Data_测试数据。报告只引用49个原始CSV，不使用历史PNG、FIG、"
        "summary或日志作为计算输入。",
    )
    add_text(
        document,
        "AD9245分析覆盖SFDR、X3G输入频率响应、X3G激励隔离度和X3G输入功率响应；"
        "AD766分析覆盖X9斜坡输出线性特性及两个直流码点的输出噪声。没有原始数据支撑的"
        "非线性、温漂、相噪和隔离度等项目按细则顺序保留，并标记为未测试。",
    )

    add_heading(document, "结论摘要", 2)
    summary_rows = [
        ["AD9245", "采样率", "≥20 MSPS", "MATLAB按25 MSPS处理；CSV无独立时间列", "暂不能判定", "配置值达到要求，缺独立测量"],
        ["AD9245", "SFDR", ">65 dB", "10 MHz时X1G/X2G/X3G为64.660/63.014/64.914 dB", "不满足", "其余名义6 dBm测点满足"],
        ["AD9245", "输入频率范围", "DC～10 MHz", "X3G −3 dB带宽3.799983 MHz", "不满足", "5 MHz和一份10 MHz文件拟合异常"],
        ["AD9245", "隔离度", "≥40 dB", "X3G→X1G/X2G/X4G为106.003/102.835/109.479 dB", "满足", "仅覆盖X3G激励工况"],
        ["AD9245", "射频路饱和输入", "5～6 dBm", "5 dBm和6 dBm文件完全相同", "暂不能判定", "6 dBm需重新采集"],
        ["AD9245", "ADC非线性/噪声/温漂", "按细则", "没有对应原始数据", "未测试", "需补测"],
        ["AD766", "输出噪声（6V码点）", "<12 µV/√Hz@1 Hz；积分<1 mVrms", "396.149 µV/√Hz；6.504 mVrms", "不满足", "采样率由时间列计算"],
        ["AD766", "输出噪声（8000码点）", "同上", "9.221 µV/√Hz；0.270 mVrms", "满足", "测试配置与负载记录不完整"],
        ["AD766", "输出范围/线性", "按细则输出范围", "X9斜坡−0.179～5.974 V；最大拟合残差79.085 mV", "暂不能判定", "缺码值映射、负载和独立线性限值"],
        ["AD766", "相噪/隔离度", "按细则", "没有对应原始数据", "未测试", "需补测"],
    ]
    add_caption(document, "表 2 核心测试结果与结论")
    add_table(
        document,
        ["对象", "项目", "指标", "实测或计算结果", "结论", "说明"],
        summary_rows,
        widths=[1.8, 2.5, 3.0, 5.2, 1.8, 3.0],
        font_size=7.7,
    )
    add_text(
        document,
        "总体结论：当前数据不能支持整板全部指标满足。AD9245在10 MHz处有三个通道SFDR低于65 dB，"
        "X3G的−3 dB带宽为3.799983 MHz，未覆盖DC～10 MHz要求；X3G激励工况下三路隔离度满足40 dB。"
        "AD766的6 V直流码点噪声明显超过限值，8000码点数值满足。输入功率5～6 dBm边界因文件重复暂不能判定。",
        bold_prefix="总体结论：",
    )

    add_page_break(document)
    add_heading(document, "测试对象、接口与公共条件", 1)
    add_heading(document, "被测对象与数据覆盖", 2)
    coverage_rows = [
        ["AD9245 X1G", "SFDR", "5", "1 MHz、5 MHz、7.5 MHz、10 MHz；另含1 MHz/2 Vpp"],
        ["AD9245 X2G", "SFDR", "4", "1 MHz、5 MHz、7.5 MHz、10 MHz"],
        ["AD9245 X3G", "SFDR", "4", "1 MHz、5 MHz、7.5 MHz、10 MHz"],
        ["AD9245 X4G", "SFDR", "4", "1 MHz、5 MHz、7.5 MHz、10 MHz"],
        ["AD9245 X3G", "输入频率响应", "12", "1 kHz～10 MHz；2个文件拟合异常"],
        ["AD9245 X3G", "隔离度", "4", "X3G输入1 MHz、7 dBm，其余三路接地"],
        ["AD9245 X3G", "输入功率", "13", "1 MHz，−10～10.8 dBm"],
        ["AD766 X9", "斜坡输出线性", "1", "示波器电压记录"],
        ["AD766", "直流输出噪声", "2", "7FFF约6 V、8000码点"],
    ]
    add_caption(document, "表 3 原始数据覆盖情况")
    add_table(document, ["对象", "测试项目", "CSV数量", "覆盖说明"], coverage_rows, widths=[3.0, 3.3, 2.0, 8.0], font_size=8.5)

    add_heading(document, "公共处理条件", 2)
    condition_rows = [
        ["AD9245采样率", "25 MHz", "四个独立分析脚本统一配置"],
        ["AD9245码制", "14 bit有符号码", "满量程峰值8192 code"],
        ["SFDR FFT", "NFFT=131072；Hann窗", "排除直流及基波积分区后搜索最大杂散"],
        ["正弦拟合", "20周期；最少1024点", "Code_pp=2×拟合振幅"],
        ["频响有效点", "R²≥0.99", "异常点保留但不参与带宽插值"],
        ["隔离度频率", "1 MHz", "20log10(激励Code_pp/安静Code_pp)"],
        ["功率标定范围", "−10～6 dBm", "重复、平台或削顶点不参与线性标定"],
        ["AD766采样率", "由CSV时间列计算", "线性文件403.226 kSPS；噪声文件416.667 kSPS"],
        ["AD766噪声谱", "4 s Hann Welch，50%重叠", "NFFT=2097152；频率分辨率0.198682 Hz"],
    ]
    add_caption(document, "表 4 数据处理公共参数")
    add_table(document, ["参数", "取值", "说明"], condition_rows, widths=[4.0, 4.2, 8.0], font_size=8.5)

    add_heading(document, "测试仪器与连接记录限制", 2)
    add_text(
        document,
        "现有CSV可确认信号频率、部分输入功率、通道名称和采样数据，但没有完整保存信号源型号、"
        "示波器型号编号、校准有效期、端接方式、线缆损耗和AD766输出负载。报告对可复算数值照实给出，"
        "涉及绝对参考面或正式验收条件时保留相应限制。",
    )

    add_page_break(document)
    add_heading(document, "测试仪器、连接和数据处理方法", 1)
    add_heading(document, "AD9245处理方法", 2)
    add_text(
        document,
        "SFDR分析先去除直流偏置，对数据加Hann窗并执行单边FFT。基波由最大有效谱线自动定位，"
        "最大杂散在排除直流和基波积分区后搜索。SFDR按基波幅度与最大杂散幅度之差计算；"
        "SNR、SINAD、THD和ENOB由同一功率谱及正弦拟合结果计算。",
    )
    add_text(
        document,
        "输入频率响应采用文件名中的频率作为已知频率进行最小二乘正弦拟合。以1 MHz及以下有效点"
        "Code_pp的中值为低频参考，按20log10(Code_pp/Code_pp,ref)计算相对幅度。−3 dB交点由相邻有效"
        "测点线性插值得到。",
    )
    add_text(
        document,
        "隔离度采用相同1 MHz频率分量的拟合幅度比。安静通道的拟合R²很低并不表示计算程序失效，"
        "而是因为总方差主要由宽带噪声构成；因此隔离度数值同时报告安静通道Code_pp、R²和残差，"
        "不能脱离噪声底单独解释。",
    )
    add_heading(document, "AD766处理方法", 2)
    add_text(
        document,
        "X9线性记录从波形前半段寻找最低点、后半段寻找最高点，截取上升斜坡后进行带截距直线拟合。"
        "残差以实测电压减拟合电压计算。因缺少每个采样点对应的DAC码值，结果只按电压偏差表示，"
        "不换算为DNL或INL LSB。",
    )
    add_text(
        document,
        "AD766噪声分析从CSV第一列时间戳计算采样率，去除平均值后采用4 s周期Hann窗、50%重叠的"
        "Welch功率谱密度估计。ASD为PSD开平方；1 kHz～100 kHz积分噪声由PSD在规定频带积分后开平方，"
        "不对ASD直接积分。",
    )

    add_page_break(document)
    add_heading(document, "四路中速AD测试", 1)
    add_heading(document, "ADC指标与覆盖", 2)
    add_text(
        document,
        "细则规定4路中速AD、采样率不低于20 MSPS、SFDR大于65 dB、通道隔离度优于40 dB、"
        "输入频率范围为DC～10 MHz，射频路饱和输入功率为5～6 dBm。本批数据能够直接计算其中"
        "SFDR、X3G频率响应、X3G激励隔离度和X3G输入功率响应。",
    )

    add_heading(document, "ADC的SFDR", 2)
    sfdr_rows = []
    for row in sorted(
        sfdr,
        key=lambda item: (
            item["Channel"],
            number(item, "FundamentalFrequencyHz"),
            item["FileName"],
        ),
    ):
        sfdr_rows.append([
            row["Channel"],
            row["FileName"],
            "{:.3f}".format(number(row, "FrequencyMHz")),
            "{:.1f}".format(number(row, "CodePp")),
            "{:.3f}".format(number(row, "SFDR")),
            "满足" if number(row, "SFDR") > 65 else "不满足",
        ])
    add_caption(document, "表 5 AD9245四通道SFDR结果")
    add_table(document, ["通道", "文件", "频率(MHz)", "Code_pp(LSB)", "SFDR(dB)", "结论"], sfdr_rows, widths=[1.5, 5.3, 2.1, 2.4, 2.0, 2.0], font_size=7.5)
    add_picture(document, FIG / "01_AD9245_SFDR_summary.png", "图 1 AD9245四通道SFDR随输入频率变化", 16.5)
    add_text(
        document,
        "1 MHz、5 MHz和7.5 MHz的名义6 dBm测点均满足65 dB要求。10 MHz时X1G、X2G和X3G分别为"
        "64.660 dB、63.014 dB和64.914 dB，低于指标；X4G为66.864 dB。X1G的1 MHz、2 Vpp文件"
        "测试条件不同，单独保留，不与6 dBm测点平均。",
    )
    add_picture(document, FIG / "02_AD9245_worst_spectra_grid.png", "图 2 各通道名义6 dBm工况下最差SFDR频谱", 16.8)
    add_text(document, "本项结论：AD9245四通道SFDR整体判定为不满足；风险集中在10 MHz高频端。", bold_prefix="本项结论：")

    add_heading(document, "ADC的输入频率测试", 2)
    bw_value = next(
        number(row, "Bandwidth3dBHz")
        for row in bandwidth
        if row["Bandwidth3dBHz"] and row["Bandwidth3dBHz"].lower() != "nan"
    )
    bw_rows = []
    for row in bandwidth:
        valid = truth(row["ValidForBandwidth"])
        relative_text = row["RelativeDb"].strip()
        bw_rows.append([
            row["FileName"],
            "{:.6g}".format(number(row, "FrequencyHz")),
            "{:.2f}".format(number(row, "CodePp")),
            "—" if not relative_text or relative_text.lower() == "nan" else "{:.3f}".format(float(relative_text)),
            "{:.6f}".format(number(row, "FitR2")),
            "有效" if valid else "剔除",
        ])
    add_caption(document, "表 6 X3G输入频率响应正弦拟合结果")
    add_table(document, ["文件", "频率(Hz)", "Code_pp(LSB)", "相对幅度(dB)", "拟合R²", "带宽计算"], bw_rows, widths=[4.0, 2.5, 2.7, 2.7, 2.4, 1.8], font_size=7.6)
    add_picture(document, FIG / "03_AD9245_bandwidth.png", "图 3 X3G Code_pp及相对幅度频率响应", 16.5)
    add_text(
        document,
        "有效测点显示幅度从1 MHz后开始明显下降。−3 dB交点为{:.6f} MHz。5MHz.ila.csv的R²约"
        "1.0×10⁻⁶，10MHz.csv的R²为0.524646，两者不参与带宽计算；另一份10Mhz_1.csv拟合有效，"
        "在10 MHz处相对幅度为−11.282 dB。".format(bw_value / 1e6),
    )
    add_text(
        document,
        "本项结论：X3G的−3 dB带宽为3.799983 MHz，低于10 MHz，判定不满足。当前数据仅覆盖X3G，"
        "其余三路频率响应未测试。",
        bold_prefix="本项结论：",
    )

    add_heading(document, "ADC的非线性测试", 2)
    add_text(document, "本批目录没有AD9245直方图法DNL/INL原始记录，不能评价INL≤20 LSB、DNL≤5 LSB。")
    add_text(document, "本项结论：未测试。", bold_prefix="本项结论：")

    add_heading(document, "ADC的隔离度测试", 2)
    isolation_rows = []
    for row in isolation:
        isolation_rows.append([
            "{}→{}".format(row["DrivenChannel"], row["QuietChannel"]),
            "{:.3f}".format(number(row, "DrivenCodePp")),
            "{:.6f}".format(number(row, "QuietCodePp")),
            "{:.3f}".format(number(row, "IsolationDb")),
            "{:.6f}".format(number(row, "QuietFitR2")),
            status_text(truth(row["Pass"])),
        ])
    add_caption(document, "表 7 X3G激励时的AD9245通道隔离度")
    add_table(document, ["通道关系", "激励Code_pp", "安静Code_pp", "隔离度(dB)", "安静通道R²", "结论"], isolation_rows, widths=[2.5, 2.6, 2.8, 2.5, 2.5, 2.0], font_size=8)
    add_picture(document, FIG / "04_AD9245_isolation.png", "图 4 X3G激励工况下三路隔离度", 15.5)
    add_text(
        document,
        "最差隔离度为X3G→X2G的102.835 dB，三路均高于40 dB。安静通道拟合分量只有"
        "0.0336～0.0722 LSB，R²接近零，说明数值已接近本批记录的测量噪声底；报告保留原始数值，"
        "不把109 dB等结果解释为系统能够稳定验证到相同动态范围。",
    )
    add_text(document, "本项结论：X3G激励工况满足；完整4×4隔离度矩阵未测试。", bold_prefix="本项结论：")

    add_heading(document, "ADC的噪声测试", 2)
    add_text(document, "当前目录没有按AD9245输入短路条件采集的噪声原始数据，不能评价≤10 µV/√Hz@1 Hz指标。")
    add_text(document, "本项结论：未测试。", bold_prefix="本项结论：")

    add_heading(document, "ADC的温漂测试", 2)
    add_text(document, "当前目录没有−10～45 ℃温度扫描数据。")
    add_text(document, "本项结论：未测试。", bold_prefix="本项结论：")

    add_heading(document, "ADC的输入功率测试", 2)
    power_rows = []
    for row in sorted(power, key=lambda item: number(item, "InputPowerDbm")):
        power_rows.append([
            "{:.1f}".format(number(row, "InputPowerDbm")),
            "{:.1f}".format(number(row, "CodePp")),
            "{:.3f}".format(number(row, "CodeRmsDbfs")),
            "{:.7f}".format(number(row, "FitR2")),
            "是" if truth(row["ClippingFlag"]) else "否",
            "是" if truth(row["PlateauFlag"]) else "否",
            "是" if truth(row["CalibrationIncluded"]) else "否",
        ])
    add_caption(document, "表 8 X3G输入功率响应结果")
    add_table(document, ["输入(dBm)", "Code_pp(LSB)", "RMS(dBFS)", "拟合R²", "削顶", "重复/平台", "纳入标定"], power_rows, widths=[2.1, 2.7, 2.5, 2.6, 1.7, 2.3, 2.0], font_size=7.7)
    add_picture(document, FIG / "05_AD9245_power_scale.png", "图 5 X3G输入功率Code_pp及dBFS响应", 16.5)
    add_text(
        document,
        "排除重复/平台点后，−10～6 dBm范围内的标定斜率为0.998387 dB/dBm，截距为"
        "−14.306572 dBFS，R²=0.99999435，说明有效点的对数幅度响应接近1 dB/dB。10.8 dBm文件"
        "有4.0001%的样本接近满量程，被识别为削顶点。",
    )
    add_text(
        document,
        "−2dBm.ila.csv与−5dBm.csv内容完全相同，6dBm.csv与5dBm.csv内容完全相同。"
        "因此5～6 dBm之间的幅度平台不能作为真实饱和证据，必须重新采集6 dBm及相邻点。",
    )
    add_text(document, "本项结论：有效点线性良好；5～6 dBm饱和输入指标暂不能判定。", bold_prefix="本项结论：")

    add_page_break(document)
    add_heading(document, "AD766测试", 1)
    add_heading(document, "指标与数据覆盖", 2)
    add_text(
        document,
        "细则要求中速DAC分辨率不低于14 bit、速率不低于300 kSPS，直流输出噪声小于"
        "12 µV/√Hz@1 Hz，1 kHz～100 kHz积分噪声小于1 mVrms，并规定不同通道的输出电压范围。"
        "本批数据仅包括X9斜坡输出和两个直流码点噪声记录。",
    )

    add_heading(document, "中速DAC相噪测试", 2)
    add_text(document, "当前目录没有1 MHz正弦输出相噪数据。")
    add_text(document, "本项结论：未测试。", bold_prefix="本项结论：")

    add_heading(document, "中速DAC噪声测试", 2)
    noise_rows = []
    for row in noise:
        noise_rows.append([
            row["fileName"],
            "{:.6f}".format(number(row, "meanVoltageV")),
            "{:.3f}".format(number(row, "noiseRmsV") * 1e3),
            "{:.3f}".format(number(row, "asdAt1HzVPerSqrtHz") * 1e6),
            "{:.3f}".format(number(row, "integratedNoise1kTo100kVrms") * 1e3),
            status_text(truth(row["asdPass"]) and truth(row["integratedPass"])),
        ])
    add_caption(document, "表 9 AD766直流输出噪声结果")
    add_table(document, ["文件", "平均电压(V)", "全带RMS(mV)", "ASD@1Hz(µV/√Hz)", "1k～100k积分(mVrms)", "数值比较"], noise_rows, widths=[4.0, 2.3, 2.5, 3.3, 3.5, 2.0], font_size=7.8)
    add_picture(document, FIG / "07_AD766_noise_6V_7FFF_DC.png", "图 6 AD766 7FFF约6 V码点输出噪声ASD", 16.2)
    add_picture(document, FIG / "08_AD766_noise_CODE_8000_DC.png", "图 7 AD766 8000码点输出噪声ASD", 16.2)
    add_text(
        document,
        "两份记录均包含10000000点、时长24.000 s，时间列给出的采样率为416.667 kSPS。"
        "7FFF约6 V码点的ASD@1 Hz为396.149 µV/√Hz，1 kHz～100 kHz积分噪声为6.504 mVrms，"
        "均超过限值。8000码点分别为9.221 µV/√Hz和0.270 mVrms，数值满足。",
    )
    add_text(
        document,
        "本项结论：在当前记录条件下，AD766输出噪声整体不满足；其中8000码点满足，7FFF约6 V码点不满足。"
        "正式复测需补充输出负载、示波器带宽、耦合方式和本底噪声。",
        bold_prefix="本项结论：",
    )

    add_heading(document, "中速DAC输出电压与X9线性", 2)
    line = linearity[0]
    line_rows = [[
        line["FileName"],
        "{:.3f}".format(number(line, "SampleRateHz") / 1e3),
        "{:.6f}".format(number(line, "RampMinimumV")),
        "{:.6f}".format(number(line, "RampMaximumV")),
        "{:.8f}".format(number(line, "FitR2")),
        "{:.3f}".format(number(line, "MaximumAbsoluteResidualV") * 1e3),
        "{:.3f}".format(number(line, "ResidualPeakToPeakV") * 1e3),
    ]]
    add_caption(document, "表 10 AD766 X9斜坡输出线性拟合")
    add_table(document, ["文件", "采样率(kSPS)", "最小值(V)", "最大值(V)", "拟合R²", "最大|残差|(mV)", "残差峰峰值(mV)"], line_rows, widths=[3.8, 2.3, 2.2, 2.2, 2.2, 2.8, 2.8], font_size=7.7)
    add_picture(document, FIG / "06_AD766_X9_linearity.png", "图 8 AD766 X9上升斜坡拟合及电压残差", 16.4)
    add_text(
        document,
        "X9上升斜坡范围为−0.178711～5.973630 V，直线拟合R²=0.99989，最大绝对残差79.085 mV，"
        "残差峰峰值154.994 mV。由于没有同步保存DAC码值序列、步进速率和负载，本结果不能换算为LSB，"
        "也不能据此验证全部通道输出范围和1 mV调节精度。",
    )
    add_text(document, "本项结论：暂不能判定；现有结果作为X9斜坡输出特性记录。", bold_prefix="本项结论：")

    add_heading(document, "中速DAC隔离度测试", 2)
    add_text(document, "当前目录没有单路DAC输出、其余通道静默条件下的多通道同步记录。")
    add_text(document, "本项结论：未测试。", bold_prefix="本项结论：")

    add_page_break(document)
    add_heading(document, "异常现象、证据限制与补测要求", 1)
    anomaly_rows = [[r["DataOrItem"], r["Observation"], r["Disposition"]] for r in anomalies]
    add_caption(document, "表 11 关键异常及处理")
    add_table(document, ["数据或项目", "异常现象", "处理及影响"], anomaly_rows, widths=[5.0, 5.5, 6.2], font_size=8)
    add_text(
        document,
        "重复文件采用SHA-256确认，不是仅凭拟合结果相同推断。异常频响文件保留在原始清单中，"
        "但不参与−3 dB插值。AD766噪声采样率由时间列计算，避免沿用原脚本4.17 MHz固定值造成"
        "频率轴约10倍偏差。",
    )

    follow_rows = [
        ["AD9245 10 MHz SFDR", "X1G/X2G/X3G低于65 dB", "保持同一输入参考面，每路重复3次，保存信号源和板端幅度", "P0"],
        ["AD9245频率响应", "仅X3G且−3 dB为3.80 MHz", "四路1 kHz～12 MHz扫频，3～5 MHz步进≤0.25 MHz", "P0"],
        ["AD9245输入功率", "6 dBm文件与5 dBm重复", "重新采集4～8 dBm，步进0.5 dB，并同步记录削顶和THD", "P0"],
        ["AD9245隔离度", "只覆盖X3G激励", "分别激励X1G～X4G，形成完整4×4矩阵", "P1"],
        ["AD9245噪声/温漂/INL-DNL", "无原始数据", "按细则补齐输入短路噪声、温度扫描和直方图法数据", "P1"],
        ["AD766 6 V噪声", "明显超限", "固定负载和示波器带宽，采集仪器本底并重复测试", "P0"],
        ["AD766输出范围与线性", "缺码值和负载记录", "同步保存DAC码值、输出电压、负载、步进速率及全部通道", "P1"],
        ["AD766相噪/隔离度", "无原始数据", "按细则补测并保存仪器原始文件", "P1"],
    ]
    add_caption(document, "表 12 后续测试计划")
    add_table(document, ["项目", "当前缺口", "补测要求", "优先级"], follow_rows, widths=[4.0, 4.5, 7.3, 1.5], font_size=8)

    add_page_break(document)
    add_heading(document, "指标实测对比总表", 1)
    compliance_rows = summary_rows
    add_caption(document, "表 13 适用指标、实测结果和正式状态")
    add_table(document, ["对象", "项目", "指标", "实测或计算结果", "结论", "说明"], compliance_rows, widths=[1.8, 2.5, 3.0, 5.2, 1.8, 3.0], font_size=7.7)
    add_text(
        document,
        "本报告结论只适用于当前文件、当前配置和已记录工况。对没有原始数据、测试条件不完整或"
        "文件重复的项目，不以趋势外推替代正式试验。",
    )

    add_page_break(document)
    add_heading(document, "附录", 1)
    add_heading(document, "计算公式与处理参数", 2)
    formula_rows = [
        ["正弦峰峰值", "Code_pp=2|A|", "A为最小二乘正弦拟合峰值，单位LSB"],
        ["相对幅度", "Gain_dB=20log10(Code_pp/Code_pp,ref)", "用于输入频率响应"],
        ["SFDR", "A_fundamental−A_largest_spur", "单位dB"],
        ["隔离度", "20log10(Code_pp,driven/Code_pp,quiet)", "单位dB，同频拟合幅度比"],
        ["ADC RMS dBFS", "20log10(Code_rms/8192)", "14 bit有符号满量程峰值8192 code"],
        ["ASD", "ASD(f)=sqrt(PSD(f))", "单位V/√Hz"],
        ["积分噪声", "Vrms=sqrt(∫PSD(f)df)", "AD766积分范围1 kHz～100 kHz"],
    ]
    add_table(document, ["项目", "公式", "说明"], formula_rows, widths=[3.0, 6.0, 7.2], font_size=8.5)

    add_heading(document, "数据与脚本索引", 2)
    index_rows = [
        ["D-SFDR", "AD9245四通道SFDR，17个CSV", "adc_sfdr_analysis.m", "SFDR"],
        ["D-BW", "X3G频率响应，12个CSV", "adc_bandwidth_analysis.m", "Bandwidth"],
        ["D-ISO", "X3G激励隔离度，4个CSV", "adc_isolation_analysis.m", "Isolation"],
        ["D-POWER", "X3G输入功率，13个CSV", "adc_power_scale_analysis.m", "PowerScale"],
        ["D-DAC-LIN", "AD766 X9斜坡，1个CSV", "run_ad766_full_analysis.m", "AD766_Linearity"],
        ["D-DAC-N", "AD766噪声，2个CSV", "run_ad766_full_analysis.m", "AD766_Noise"],
    ]
    add_caption(document, "表 14 数据、脚本与结果目录索引")
    add_table(document, ["编号", "内容", "处理脚本", "结果目录"], index_rows, widths=[2.2, 5.5, 5.0, 3.5], font_size=8.5)
    add_text(document, "MATLAB程序根目录：C:\\Users\\86183\\Desktop\\00_matlab。")
    add_text(document, "本次复算结果包：{}。".format(str(BUNDLE)))
    add_text(document, "源文件清单：source_manifest.csv，共{}项；每项记录相对路径、字节数和SHA-256。".format(len(manifest)))
    add_text(document, "运行日志：Logs\\AD9245_matlab_run.log、Logs\\AD766_matlab_run.log。")

    add_heading(document, "源文件数量核对", 2)
    count_rows = []
    counts = {}
    for row in manifest:
        prefix = row["RelativePath"].split("\\")[0]
        if prefix not in counts:
            counts[prefix] = [0, 0]
        counts[prefix][0] += 1
        counts[prefix][1] += int(row["SizeBytes"])
    for prefix in sorted(counts):
        count_rows.append([prefix, counts[prefix][0], counts[prefix][1]])
    add_table(document, ["数据类别", "CSV数量", "总字节数"], count_rows, widths=[6.0, 3.0, 6.0], font_size=9)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT_DOCX))
    print(str(OUTPUT_DOCX))


if __name__ == "__main__":
    build_report()
