from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"F:\01_Laser\20260727_513test\02_GS与延迟驱动测试")
DATA = ROOT / r"02_Data_测试数据\513_GS_AD_DATA\AD9245"
REPORT_DIR = ROOT / r"01_Documents_文档\02_Test_Reports_测试报告"
OUTPUT = REPORT_DIR / "AD9245_四通道测试结果与补测建议_20260802.docx"

CHANNELS = ["X1G", "X2G", "X3G", "X4G"]
NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "EEF5FA"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
LIGHT_YELLOW = "FFF2CC"
LIGHT_GRAY = "E7E6E6"
TEXT = "222222"


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def num(value, default=math.nan):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def fmt(value, digits=3):
    value = num(value)
    if math.isnan(value):
        return "未记录"
    return f"{value:.{digits}f}".replace("-", "−", 1) if value < 0 else f"{value:.{digits}f}"


def fmt_db(value):
    value = num(value)
    if math.isnan(value):
        return "未记录"
    return fmt(value, 3) + " dB"


def fmt_range(values, digits=3, unit=""):
    values = [num(v) for v in values if not math.isnan(num(v))]
    if not values:
        return "未记录"
    lo, hi = min(values), max(values)
    lo_text = fmt(lo, digits)
    hi_text = fmt(hi, digits)
    if abs(hi - lo) < 10 ** (-(digits + 1)):
        return lo_text + unit
    return lo_text + "～" + hi_text + unit


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_font(run, name="宋体", size=10, bold=False, color=TEXT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.right_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    return p


def write_cell(cell, text, bold=False, size=9, fill=None, color=TEXT, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = clear_cell(cell)
    p.alignment = align
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if fill:
        set_cell_shading(cell, fill)


def status_fill(status):
    return {
        "满足": LIGHT_GREEN,
        "不满足": LIGHT_RED,
        "暂不能判定": LIGHT_YELLOW,
        "未测试": LIGHT_GRAY,
    }.get(status, None)


def add_table(doc, headers, rows, widths, font_size=8.5, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed_layout(table)
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_width(cell, widths[i])
        write_cell(cell, header, bold=True, size=font_size, fill=NAVY, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for i, (cell, value) in enumerate(zip(row.cells, row_values)):
            set_cell_width(cell, widths[i])
            fill = None
            if status_col is not None and i == status_col:
                fill = status_fill(str(value))
            write_cell(cell, value, size=font_size, fill=fill, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_font(run2, size=9, color="666666")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, name="黑体", size=15 if level == 1 else 11.5, bold=True, color=NAVY)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=10.5, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, size=10.5)
    else:
        r = p.add_run(text)
        set_font(r, size=10.5)
    return p


def parse_frequency_from_name(name):
    m = re.search(r"_(\d+(?:\.\d+)?)MHz", name, re.I)
    return float(m.group(1)) if m else math.nan


def load_results():
    sfdr = {ch: read_csv(DATA / ch / "SFDR" / "results" / "ADC_SFDR_summary.csv") for ch in CHANNELS}
    bw = {ch: read_csv(DATA / ch / "FrequencyResponse" / "results" / "ADC_bandwidth_summary.csv") for ch in CHANNELS}
    power = {ch: read_csv(DATA / ch / "InputPowerScale" / "results" / "ADC_power_scale_summary.csv") for ch in CHANNELS}
    iso = {ch: read_csv(DATA / ch / "Isolation" / "results" / "ADC_isolation_summary.csv") for ch in CHANNELS}
    inldnl = {ch: read_csv(DATA / ch / "INL_DNL" / "results" / "ADC_inl_dnl_summary.csv")[0] for ch in CHANNELS}
    return sfdr, bw, power, iso, inldnl


def sfdr_rows(sfdr):
    rows = []
    for ch in CHANNELS:
        grouped = defaultdict(list)
        for r in sfdr[ch]:
            m = re.search(r"_(\d+(?:\.\d+)?)MHz", r["FileName"], re.I)
            key = float(m.group(1)) if m else math.nan
            grouped[key].append(num(r["SFDR"]))
        cells = []
        for f in (1.0, 5.0, 7.5, 10.0):
            cells.append(fmt_range(grouped.get(f, []), 3, ""))
        minimum = min(v for vals in grouped.values() for v in vals)
        status = "满足" if minimum > 65.0 else "不满足"
        rows.append([ch, *cells, f"{minimum:.3f} dB", status])
    return rows


def bandwidth_rows(bw):
    rows = []
    for ch in CHANNELS:
        all_rows = bw[ch]
        valid = [r for r in all_rows if num(r["ValidForBandwidth"]) == 1]
        ten = [num(r["RelativeDb"]) for r in valid if abs(parse_frequency_from_name(r["FileName"]) - 10.0) < 0.01]
        invalid = [r["FileName"] for r in all_rows if num(r["ValidForBandwidth"]) != 1]
        rows.append([
            ch,
            fmt(num(all_rows[0]["Bandwidth3dBHz"]) / 1e6, 3) + " MHz",
            f"{len(valid)}/{len(all_rows)}",
            fmt_range(ten, 3, " dB"),
            "；".join(invalid[:3]) if invalid else "无",
            "暂不能判定",
        ])
    return rows


def inldnl_rows(inldnl):
    rows = []
    for ch in CHANNELS:
        r = inldnl[ch]
        status = "满足" if num(r["MaxAbsDNL_LSB"]) <= 5 and num(r["MaxAbsINL_LSB"]) <= 20 else "不满足"
        rows.append([
            ch,
            r["FileCount"],
            fmt(num(r["FrequencyHz"]) / 1e6, 6) + " MHz",
            fmt(r["MaxAbsDNL_LSB"]) + " LSB",
            fmt(r["MaxAbsINL_LSB"]) + " LSB",
            status,
        ])
    return rows


def isolation_rows(iso):
    matrix = {d: {q: None for q in CHANNELS} for d in CHANNELS}
    worst_rows = []
    for d in CHANNELS:
        vals = []
        for r in iso[d]:
            q = r["QuietChannel"]
            val = num(r["IsolationDb"])
            matrix[d][q] = val
            vals.append((val, q))
        worst, q = min(vals)
        worst_rows.append([d, q, f"{worst:.3f} dB", "满足" if worst > 40 else "不满足"])
    rows = []
    for d in CHANNELS:
        rows.append([d] + ["-" if matrix[d][q] is None else f"{matrix[d][q]:.3f}" for q in CHANNELS])
    return rows, worst_rows


def power_rows(power):
    rows = []
    for ch in CHANNELS:
        rs = power[ch]
        first = rs[0]
        p5 = [r for r in rs if abs(num(r["InputPowerDbm"]) - 5) < 0.01]
        p6 = [r for r in rs if abs(num(r["InputPowerDbm"]) - 6) < 0.01]
        notes = []
        if any(num(r["ClippingFlag"]) for r in rs):
            notes.append("存在削顶标记")
        if sum("run02" in r["FileName"] for r in rs):
            notes.append("存在重复采集")
        rows.append([
            ch,
            f"{num(first['CalibrationSlopeDbPerDbm']):.3f}",
            f"{num(first['CalibrationInterceptDb']):.3f}",
            f"{num(first['CalibrationR2']):.6f}",
            fmt_range([r["CodePp"] for r in p5], 1, ""),
            fmt_range([r["CodePp"] for r in p6], 1, ""),
            "；".join(notes) if notes else "无",
            "暂不能判定",
        ])
    return rows


def make_doc():
    sfdr, bw, power, iso, inldnl = load_results()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
        st = doc.styles[name]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("GS延迟驱动板  |  AD9245四通道测试简报")
    set_font(hr, size=8.5, color="666666")
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("AD9245四通道测试结果与补测建议")
    set_font(r, name="黑体", size=23, bold=True, color=NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(10)
    r = p2.add_run("简要报告  |  2026年8月2日")
    set_font(r, size=11, color="666666")

    info_rows = [
        ["测试对象", "YCQD板 AD9245，通道 X1G～X4G", "数据目录", r"02_Data_测试数据\513_GS_AD_DATA\AD9245"],
        ["处理配置", "14 bit；signed；采样率配置 20 MHz；满量程峰值码 8192", "结果依据", "当前目录下各测试项目 results 汇总 CSV"],
        ["报告用途", "对照测试细则做初步判断，列出最终判定所需的补测数据", "原始数据", "未修改；结果图保留在各测试项目 results 目录"],
    ]
    add_table(doc, ["项目", "内容", "项目", "内容"], info_rows, [1.1, 3.75, 1.1, 4.3], font_size=9)

    add_heading(doc, "1 测试目的与结论摘要", 1)
    add_body(doc, "现有结果中，四通道 INL/DNL 和通道隔离度达到测试细则限值。SFDR 在 X1G、X2G、X3G 的 10 MHz 测点低于 65 dB，X4G 的已测点高于 65 dB。")
    add_body(doc, "频率响应、输入功率饱和边界和 ADC 实际采样率还缺少最终判定所需条件；ADC 噪声和温漂没有对应原始数据。下表的结论只针对当前数据覆盖范围。")

    summary_rows = [
        ["AD9245", "采样率", "≥20 MSPS", "配置20 MHz；需外部实测", "X1G～X4G：20 MHz为处理配置值", "暂不能判定", "CSV无独立时标"],
        ["AD9245", "SFDR", ">65 dB；1/5/10 MHz", "1、5、7.5、10 MHz；名义6 dBm", "X1G/X2G/X3G最低值64.660/63.014/64.914 dB；X4G为66.864 dB", "不满足", "X1G～X3G在10 MHz低于65 dB"],
        ["AD9245", "频率响应/3 dB带宽", "DC～10 MHz；细则未给带宽下限", "正弦拟合；低频有效点归一化", "X1G/X2G/X3G/X4G：4.893/5.037/3.864/3.966 MHz", "暂不能判定", "10 MHz相对幅度为−8.390～−11.279 dB"],
        ["AD9245", "INL/DNL", "INL≤20 LSB；DNL≤5 LSB", "每通道60组；码密度统计", "最大|DNL|0.248～0.304 LSB；最大|INL|5.140～8.709 LSB", "满足", "四通道均在限值内"],
        ["AD9245", "隔离度", ">40 dB", "1 MHz、7 dBm；一路激励，其余安静", "12组关系均高于40 dB；整体最差92.604 dB", "满足", "安静通道R²低，不作为剔除条件"],
        ["AD9245", "输入功率响应", "1 MHz；−10～6 dBm；核查5～6 dBm饱和", "Code_rms_dBFS与输入功率拟合", "斜率0.999～1.010 dB/dBm；R²0.967694～0.999995", "暂不能判定", "X3G有重复采集；饱和边界缺少重复点"],
        ["AD9245", "ADC噪声", "≤10 µV/√Hz @1 Hz", "四通道输入接地或50 Ω终端", "X1G～X4G：未测试", "未测试", "无满足1 Hz条件的原始数据"],
        ["AD9245", "温漂", "≤1 mV/℃；−10～45 ℃，5 ℃步进", "四通道输入接地；3个循环", "X1G～X4G：未测试", "未测试", "无温度扫描数据"],
    ]
    add_table(doc, ["对象", "测试项目", "指标", "测试条件", "实测结果", "结论", "备注"], summary_rows, [0.75, 1.15, 1.75, 2.0, 3.15, 1.0, 1.55], font_size=7.7, status_col=5)

    add_heading(doc, "2 测试对象、系统连接与测试条件", 1)
    add_body(doc, "被测对象为 YCQD 板 AD9245 四路输入通道 X1G、X2G、X3G 和 X4G。SFDR、频率响应、INL/DNL 和输入功率项目按各通道分别采集；隔离度项目按一路激励、其余通道安静的方式记录。")
    add_body(doc, "结果文件位于 AD9245 各通道下的 SFDR、FrequencyResponse、INL_DNL、Isolation 和 InputPowerScale\\results 目录。报告采用现有汇总 CSV，不重新解释旧报告中的结果。")
    condition_rows = [
        ["ADC配置", "14 bit signed，满量程峰值码8192", "来自分析脚本配置；用于dBFS和削顶判断"],
        ["采样率", "20 MHz", "处理配置值；不是AD9245采样时钟的独立实测值"],
        ["SFDR处理", "Hann窗FFT；自动搜索基波和最大杂散", "SFDR = 基波幅度 − 最大杂散幅度"],
        ["频率响应", "已知频率正弦拟合；低频有效点作参考", "Code_pp = 2|A|；相对幅度 = 20log10(Code_pp/Code_pp,ref)"],
        ["INL/DNL", "每通道60组数据；码密度统计", "按汇总文件中的MaxAbsDNL和MaxAbsINL比较"],
        ["隔离度", "1 MHz、7 dBm；激励一路，其余通道安静", "Isolation = 20log10(DrivenCode_pp/QuietCode_pp)"],
        ["功率响应", "1 MHz；输入功率−10～6 dBm", "Code_rms_dBFS与输入功率线性拟合"],
    ]
    add_table(doc, ["项目", "处理条件", "说明"], condition_rows, [1.25, 3.7, 4.4], font_size=8.5)

    add_heading(doc, "3 分项测试结果", 1)
    add_heading(doc, "3.1 SFDR", 2)
    add_body(doc, "测试结果覆盖 1 MHz、5 MHz、7.5 MHz 和 10 MHz。X1G 的 1 MHz有两次采集，表中给出范围；其余频点按结果文件列出。10 MHz测点是当前判定的限制点。")
    add_table(doc, ["通道", "1 MHz SFDR (dB)", "5 MHz SFDR (dB)", "7.5 MHz SFDR (dB)", "10 MHz SFDR (dB)", "最低值 (dB)", "结论"], sfdr_rows(sfdr), [0.85, 1.45, 1.45, 1.55, 1.55, 1.2, 1.0], font_size=8.5, status_col=6)

    add_heading(doc, "3.2 输入频率响应和3 dB带宽", 2)
    add_body(doc, "带宽由低频有效测点归一化后插值获得。X1G～X4G的3 dB交点为3.864～5.037 MHz。测试细则给出DC～10 MHz输入范围，但没有给出最低3 dB带宽值；因此本项暂不能判定。")
    add_table(doc, ["通道", "3 dB带宽", "有效点/总点", "10 MHz相对幅度", "异常或剔除文件", "结论"], bandwidth_rows(bw), [0.8, 1.25, 1.0, 1.45, 4.35, 1.0], font_size=8.2, status_col=5)
    add_body(doc, "频率文件存在命名和实际频率不一致：各通道1 kHz文件实际约1.022 kHz且拟合R²低于0.99；X1G的5 MHz文件实际约10 MHz；X3G的5 MHz文件实际约4 MHz；X3G另有一个10 MHz文件R²约0.525。上述文件没有删除，带宽计算仅使用有效点。")

    add_heading(doc, "3.3 INL/DNL", 2)
    add_body(doc, "四通道均采用60组数据，实测频率约1.012335 MHz。最大DNL为0.248～0.304 LSB，最大INL为5.140～8.709 LSB，均低于测试细则限值。")
    add_table(doc, ["通道", "数据组数", "拟合频率", "最大|DNL|", "最大|INL|", "结论"], inldnl_rows(inldnl), [0.9, 1.1, 1.65, 1.55, 1.55, 1.0], font_size=8.8, status_col=5)

    doc.add_page_break()
    add_heading(doc, "3.4 通道隔离度", 2)
    add_body(doc, "隔离度按激励通道与安静通道在1 MHz处的Code_pp比值计算。12组通道关系均高于40 dB，整体最差值为X1G激励、X4G安静时的92.604 dB。安静通道拟合R²较低属于噪声底条件下的正常现象，不作为剔除条件。")
    matrix_rows, worst_rows = isolation_rows(iso)
    add_table(doc, ["激励\\安静", "X1G", "X2G", "X3G", "X4G"], matrix_rows, [1.5, 1.9, 1.9, 1.9, 1.9], font_size=9)
    add_table(doc, ["激励通道", "最差安静通道", "最差隔离度", "结论"], worst_rows, [1.5, 2.0, 2.0, 1.2], font_size=9, status_col=3)

    add_heading(doc, "3.5 输入功率响应", 2)
    add_body(doc, "四通道1 MHz输入功率响应的拟合斜率接近1 dB/dBm。X1G、X2G和X4G的拟合R²约0.999995；X3G的R²为0.967694，原因与重复文件和−2 dBm、6 dBm处数据不一致有关。现有数据没有形成5～6 dBm饱和边界的重复测量，功率响应项目暂不能判定。")
    add_table(doc, ["通道", "标定斜率 (dB/dBm)", "截距 (dBFS)", "R²", "5 dBm Code_pp", "6 dBm Code_pp", "数据情况", "结论"], power_rows(power), [0.8, 1.45, 1.2, 1.0, 1.25, 1.25, 2.4, 1.0], font_size=7.8, status_col=7)

    add_heading(doc, "4 异常现象与测试限制", 1)
    add_body(doc, "频率文件存在命名偏差，X3G 功率数据存在重复；ADC 实际采样时钟、1 Hz 噪声和温度记录未提供。20 MHz 为本次处理配置值。")

    heading5 = add_heading(doc, "5 总结与后续测试", 1)
    heading5.paragraph_format.page_break_before = True
    follow_rows = [
        ["SFDR", "X1G～X3G的10 MHz点低于65 dB", "X1G～X4G在1/5/10 MHz各重复3次；记录输入功率、时钟和采样率", "区分偶发采集问题、输入链路影响和通道固有结果", "高"],
        ["频率响应", "1 kHz文件频率偏差；部分文件名与实际频率不一致；3 dB下限未规定", "按DC、1 kHz、100 kHz、200 kHz、500 kHz、1/2/3/5/7/10 MHz采集，每点3次；增加3～6 MHz密集点", "确认10 MHz判据并提高3 dB交点精度", "高"],
        ["INL/DNL", "当前结果满足限值，但测试条件和原始证据需归档", "保留每通道60组原始文件；记录1.01234149 MHz、输入幅度、偏置和削顶状态", "保证复测条件可追溯", "中"],
        ["隔离度", "12组结果均满足，安静通道幅度接近噪声底", "每个激励通道重复3次；记录激励/安静通道、7 dBm、终端状态和噪声底", "确认通道映射和噪声底限制", "中"],
        ["输入功率响应", "X3G存在重复文件；5～6 dBm饱和边界尚未闭合", "补采−10、−5、−2、0、1、2、3、4、5、5.5、6 dBm；5/5.5/6 dBm各3次；用功率计记录实际输入功率", "确认饱和点并排除信号源设置误差", "高"],
        ["ADC噪声", "无1 Hz判定数据", "四通道接地或50 Ω终端；延长记录使频率分辨率接近1 Hz；保存1 Hz ASD和1 kHz～100 kHz积分噪声", "验证≤10 µV/√Hz @1 Hz", "高"],
        ["温漂", "无温度扫描数据", "−10～45 ℃，5 ℃步进；每点变温5 min、稳定3 min、3个循环；四通道接地并记录平均码值", "计算mV/℃并验证≤1 mV/℃", "高"],
        ["采样率", "20 MHz为处理配置值，无AD时钟实测证据", "用示波器或频率计测AD9245实际采样时钟；同时记录FPGA/ILA时钟、占空比和时钟来源", "确认≥20 MSPS判据", "高"],
    ]
    add_table(doc, ["项目", "当前缺口或异常", "建议补充的数据", "补测目的", "优先级"], follow_rows, [1.0, 2.2, 4.1, 2.3, 0.8], font_size=7.8)

    doc.add_page_break()
    add_heading(doc, "附录 A 结果文件索引", 1)
    index_rows = [
        ["SFDR", r"AD9245\X1G～X4G\SFDR\results\ADC_SFDR_summary.csv", "adc_sfdr_analysis.m", "各通道频谱图和汇总表"],
        ["频率响应", r"AD9245\X1G～X4G\FrequencyResponse\results\ADC_bandwidth_summary.csv", "adc_bandwidth_analysis.m", "频响图和带宽汇总表"],
        ["INL/DNL", r"AD9245\X1G～X4G\INL_DNL\results\ADC_inl_dnl_summary.csv", "adc_inl_dnl_analysis.m", "码密度结果和曲线文件"],
        ["隔离度", r"AD9245\X1G～X4G\Isolation\results\ADC_isolation_summary.csv", "adc_isolation_analysis.m", "隔离度矩阵和结果图"],
        ["输入功率", r"AD9245\X1G～X4G\InputPowerScale\results\ADC_power_scale_summary.csv", "adc_power_scale_analysis.m", "功率响应图和汇总表"],
    ]
    add_table(doc, ["测试项目", "结果文件（项目相对路径）", "处理脚本", "内容"], index_rows, [1.15, 5.0, 2.2, 2.0], font_size=8.2)
    add_body(doc, "说明：本报告用于当前数据的初步对比。最终验收应以补齐测试条件、重复性数据和外部时钟/噪声测量记录后的结果为准。")

    # Keep all document properties explicit and neutral.
    doc.core_properties.title = "AD9245四通道测试结果与补测建议"
    doc.core_properties.subject = "AD9245测试结果与测试细则对比"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    make_doc()
